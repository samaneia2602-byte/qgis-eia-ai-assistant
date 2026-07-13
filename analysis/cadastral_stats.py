# -*- coding: utf-8 -*-

import os
import processing
from collections import defaultdict

from qgis.PyQt.QtCore import QVariant
from qgis.core import (
    QgsDistanceArea,
    QgsFeature,
    QgsField,
    QgsFields,
    QgsProcessingFeatureSourceDefinition,
    QgsProject,
    QgsVectorFileWriter,
    QgsVectorLayer,
    QgsWkbTypes,
)


JIMOK_SOURCE_FIELDS = (
    "지목",
    "jimok_cls",
    "jibun",
    "bonbun",
    "bubun",
    "addr",
)


def _log(callback, message):
    if callback:
        callback(message)


def _is_polygon_layer(layer):
    return (
        isinstance(layer, QgsVectorLayer)
        and layer.isValid()
        and QgsWkbTypes.geometryType(layer.wkbType())
        == QgsWkbTypes.PolygonGeometry
    )


def _field_names(layer):
    return {
        field.name().lower(): field.name()
        for field in layer.fields()
    }


def _cadastral_score(layer):
    names = _field_names(layer)
    score = 0

    for candidate in JIMOK_SOURCE_FIELDS:
        if candidate.lower() in names:
            score += 15

    name = layer.name().lower()
    for keyword in ("지적", "연속지적", "vworld", "cadastral"):
        if keyword in name:
            score += 30

    return score


def _business_score(layer):
    score = 0
    name = layer.name().lower()

    for keyword in (
        "사업", "사업지역", "경계", "부지",
        "구역", "boundary", "site", "bo",
    ):
        if keyword in name:
            score += 15

    if layer.selectedFeatureCount() > 0:
        score += 40

    if _cadastral_score(layer) == 0:
        score += 5

    return score


def find_analysis_layers(iface):
    polygon_layers = [
        layer
        for layer in QgsProject.instance().mapLayers().values()
        if _is_polygon_layer(layer)
    ]

    if len(polygon_layers) < 2:
        raise RuntimeError(
            "폴리곤 레이어가 2개 이상 필요합니다. "
            "사업지역 SHP와 연속지적도 레이어를 먼저 불러오세요."
        )

    active = iface.activeLayer()
    cadastral = sorted(
        polygon_layers,
        key=_cadastral_score,
        reverse=True,
    )[0]

    if _cadastral_score(cadastral) <= 0:
        raise RuntimeError(
            "연속지적도 레이어를 찾지 못했습니다. "
            "'지목별로 테이블 정리해줘'를 먼저 실행하세요."
        )

    business_candidates = [
        layer
        for layer in polygon_layers
        if layer.id() != cadastral.id()
    ]

    if active in business_candidates:
        business = active
    else:
        business = sorted(
            business_candidates,
            key=_business_score,
            reverse=True,
        )[0]

    return business, cadastral


def _source_definition(layer, selected_only=False):
    use_selected = selected_only and layer.selectedFeatureCount() > 0
    if use_selected:
        return QgsProcessingFeatureSourceDefinition(
            layer.id(),
            selectedFeaturesOnly=True,
        )
    return layer


def _resolve_jimok_field(layer):
    names = _field_names(layer)
    for candidate in ("지목", "jimok_cls"):
        if candidate.lower() in names:
            return names[candidate.lower()]

    raise RuntimeError(
        "연속지적도에 '지목' 필드가 없습니다. "
        "'지목별로 테이블 정리해줘'를 먼저 실행하세요."
    )


def _fix_geometries(input_layer, name, log_callback=None):
    _log(log_callback, "%s 도형 오류를 자동 수정하는 중입니다..." % name)

    fixed = processing.run(
        "native:fixgeometries",
        {
            "INPUT": input_layer,
            "OUTPUT": "memory:",
        },
    )["OUTPUT"]

    fixed.setName("%s_도형수정" % name)
    _log(
        log_callback,
        "%s 도형 수정 완료: %s개 객체"
        % (name, fixed.featureCount()),
    )
    return fixed


def _create_spatial_index(layer, name, log_callback=None):
    _log(log_callback, "%s 공간 인덱스를 생성하는 중입니다..." % name)

    try:
        processing.run(
            "native:createspatialindex",
            {"INPUT": layer},
        )
        _log(log_callback, "%s 공간 인덱스 생성 완료" % name)
    except Exception as exc:
        _log(
            log_callback,
            "%s 공간 인덱스 생성은 생략했습니다: %s"
            % (name, exc),
        )


def clip_cadastral(
    business_layer,
    cadastral_layer,
    selected_only=False,
    add_clip_layer=True,
    log_callback=None,
):
    overlay = _source_definition(
        business_layer,
        selected_only=selected_only,
    )

    if business_layer.crs() != cadastral_layer.crs():
        _log(
            log_callback,
            "사업지역 CRS를 연속지적도 CRS에 맞추는 중입니다...",
        )
        overlay = processing.run(
            "native:reprojectlayer",
            {
                "INPUT": overlay,
                "TARGET_CRS": cadastral_layer.crs(),
                "OUTPUT": "memory:",
            },
        )["OUTPUT"]
        overlay.setName("사업지역_CRS변환")
        _log(log_callback, "사업지역 CRS 변환 완료")

    fixed_overlay = _fix_geometries(
        overlay,
        "사업지역",
        log_callback,
    )
    fixed_cadastral = _fix_geometries(
        cadastral_layer,
        "연속지적도",
        log_callback,
    )

    _create_spatial_index(
        fixed_overlay,
        "사업지역",
        log_callback,
    )
    _create_spatial_index(
        fixed_cadastral,
        "연속지적도",
        log_callback,
    )

    _log(
        log_callback,
        "사업지역과 연속지적도를 중첩하여 자르는 중입니다...",
    )

    clipped = processing.run(
        "native:clip",
        {
            "INPUT": fixed_cadastral,
            "OVERLAY": fixed_overlay,
            "OUTPUT": "memory:",
        },
    )["OUTPUT"]

    _log(
        log_callback,
        "Clip 완료: %s개 객체" % clipped.featureCount(),
    )

    _log(
        log_callback,
        "멀티파트 도형을 단일파트로 정리하는 중입니다...",
    )

    singleparts = processing.run(
        "native:multiparttosingleparts",
        {
            "INPUT": clipped,
            "OUTPUT": "memory:",
        },
    )["OUTPUT"]

    singleparts.setName("사업지역_연속지적도_클립")

    if add_clip_layer:
        QgsProject.instance().addMapLayer(singleparts)

    _log(
        log_callback,
        "도형 정리 완료: %s개 객체"
        % singleparts.featureCount(),
    )

    return singleparts


def summarize_by_jimok(clipped_layer, log_callback=None):
    _log(log_callback, "지목별 면적을 계산하는 중입니다...")

    jimok_field = _resolve_jimok_field(clipped_layer)

    distance = QgsDistanceArea()
    distance.setSourceCrs(
        clipped_layer.crs(),
        QgsProject.instance().transformContext(),
    )
    distance.setEllipsoid("GRS80")

    summary = defaultdict(
        lambda: {"count": 0, "area_m2": 0.0}
    )

    for feature in clipped_layer.getFeatures():
        geometry = feature.geometry()

        if not geometry or geometry.isEmpty():
            continue

        jimok = str(feature[jimok_field] or "").strip()
        if not jimok:
            jimok = "미분류"

        area_m2 = abs(distance.measureArea(geometry))
        if area_m2 <= 0:
            continue

        summary[jimok]["count"] += 1
        summary[jimok]["area_m2"] += area_m2

    if not summary:
        raise RuntimeError(
            "사업지역과 연속지적도의 중첩 결과가 없습니다."
        )

    total_count = sum(
        item["count"] for item in summary.values()
    )
    total_area = sum(
        item["area_m2"] for item in summary.values()
    )

    rows = []
    for jimok, item in sorted(
        summary.items(),
        key=lambda pair: pair[1]["area_m2"],
        reverse=True,
    ):
        area_m2 = item["area_m2"]
        rows.append(
            {
                "지목": jimok,
                "필지수": item["count"],
                "면적_m2": area_m2,
                "면적_ha": area_m2 / 10000.0,
                "구성비_pct": (
                    area_m2 / total_area * 100.0
                    if total_area else 0.0
                ),
            }
        )

    _log(
        log_callback,
        "지목별 면적 계산 완료: %s개 지목"
        % len(rows),
    )

    return rows, total_count, total_area


def create_summary_layer(rows, total_count, total_area, add_to_project=True):
    layer = QgsVectorLayer(
        "None",
        "사업지역_지목별_면적",
        "memory",
    )
    provider = layer.dataProvider()

    fields = QgsFields()
    fields.append(QgsField("순번", QVariant.Int))
    fields.append(QgsField("지목", QVariant.String, len=30))
    fields.append(QgsField("필지수", QVariant.Int))
    fields.append(
        QgsField("면적_m2", QVariant.Double, len=20, prec=2)
    )
    fields.append(
        QgsField("면적_ha", QVariant.Double, len=20, prec=4)
    )
    fields.append(
        QgsField("구성비_pct", QVariant.Double, len=10, prec=2)
    )

    provider.addAttributes(fields)
    layer.updateFields()

    features = []
    for index, row in enumerate(rows, start=1):
        feature = QgsFeature(layer.fields())
        feature.setAttributes(
            [
                index,
                row["지목"],
                row["필지수"],
                round(row["면적_m2"], 2),
                round(row["면적_ha"], 4),
                round(row["구성비_pct"], 2),
            ]
        )
        features.append(feature)

    total_feature = QgsFeature(layer.fields())
    total_feature.setAttributes(
        [
            len(rows) + 1,
            "합계",
            total_count,
            round(total_area, 2),
            round(total_area / 10000.0, 4),
            100.0,
        ]
    )
    features.append(total_feature)

    provider.addFeatures(features)
    layer.updateExtents()

    if add_to_project:
        QgsProject.instance().addMapLayer(layer)

    return layer


def export_summary_xlsx(
    summary_layer,
    output_path,
    rows,
    total_count,
    total_area,
    report_style=True,
):
    if not output_path.lower().endswith(".xlsx"):
        output_path += ".xlsx"

    options = QgsVectorFileWriter.SaveVectorOptions()
    options.driverName = "XLSX"
    options.fileEncoding = "UTF-8"
    options.layerName = "지목별면적"
    options.actionOnExistingFile = (
        QgsVectorFileWriter.CreateOrOverwriteFile
    )

    result = QgsVectorFileWriter.writeAsVectorFormatV2(
        summary_layer,
        output_path,
        QgsProject.instance().transformContext(),
        options,
    )

    error_code = result[0] if isinstance(result, tuple) else result
    if error_code != QgsVectorFileWriter.NoError:
        error_message = (
            result[1]
            if isinstance(result, tuple) and len(result) > 1
            else ""
        )
        raise RuntimeError(
            "Excel 파일 저장에 실패했습니다. %s" % error_message
        )

    if report_style:
        try:
            _format_excel_report(
                output_path,
                rows,
                total_count,
                total_area,
            )
        except ImportError:
            pass

    return output_path


def _format_excel_report(
    output_path,
    rows,
    total_count,
    total_area,
):
    from openpyxl import load_workbook
    from openpyxl.styles import (
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )

    workbook = load_workbook(output_path)
    worksheet = workbook.active
    worksheet.title = "지목별면적"

    worksheet.insert_rows(1, amount=2)
    worksheet.merge_cells("A1:F1")
    worksheet["A1"] = "사업지역 지목별 면적현황"
    worksheet["A1"].font = Font(bold=True, size=14)
    worksheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9E2F3",
    )
    total_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9",
    )
    thin = Side(style="thin", color="808080")
    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    for cell in worksheet[3]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    max_row = worksheet.max_row
    for row in worksheet.iter_rows(
        min_row=4,
        max_row=max_row,
        min_col=1,
        max_col=6,
    ):
        for cell in row:
            cell.border = border

    for row_index in range(4, max_row + 1):
        worksheet.cell(row=row_index, column=3).number_format = "#,##0"
        worksheet.cell(row=row_index, column=4).number_format = "#,##0.00"
        worksheet.cell(row=row_index, column=5).number_format = "0.0000"
        worksheet.cell(row=row_index, column=6).number_format = "0.00"

    total_row = max_row
    for cell in worksheet[total_row]:
        cell.font = Font(bold=True)
        cell.fill = total_fill

    widths = {
        "A": 10,
        "B": 16,
        "C": 12,
        "D": 18,
        "E": 14,
        "F": 14,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    worksheet.freeze_panes = "A4"
    worksheet.sheet_view.showGridLines = False

    workbook.save(output_path)


def format_chat_table(rows, total_count, total_area):
    lines = [
        "",
        "사업지역 지목별 면적현황",
        "----------------------------------------------",
        "지목 | 필지수 | 면적(㎡) | 면적(ha) | 구성비(%)",
        "----------------------------------------------",
    ]

    for row in rows:
        lines.append(
            "%s | %s | %s | %.4f | %.2f"
            % (
                row["지목"],
                format(row["필지수"], ","),
                format(round(row["면적_m2"], 2), ",.2f"),
                row["면적_ha"],
                row["구성비_pct"],
            )
        )

    lines.extend(
        [
            "----------------------------------------------",
            "합계 | %s | %s | %.4f | 100.00"
            % (
                format(total_count, ","),
                format(round(total_area, 2), ",.2f"),
                total_area / 10000.0,
            ),
            "",
        ]
    )
    return lines


def run_cadastral_area_analysis(
    iface,
    output_path=None,
    options=None,
    log_callback=None,
):
    options = options or {}

    _log(
        log_callback,
        "사업지역 및 연속지적도 레이어를 찾는 중입니다...",
    )

    business, cadastral = find_analysis_layers(iface)

    _log(log_callback, "사업지역 레이어: %s" % business.name())
    _log(log_callback, "연속지적도 레이어: %s" % cadastral.name())

    clipped = clip_cadastral(
        business,
        cadastral,
        selected_only=options.get("selected_only", False),
        add_clip_layer=options.get("add_clip_layer", True),
        log_callback=log_callback,
    )

    rows, total_count, total_area = summarize_by_jimok(
        clipped,
        log_callback,
    )

    summary_layer = create_summary_layer(
        rows,
        total_count,
        total_area,
        add_to_project=options.get("add_summary_layer", True),
    )

    saved_path = None
    if options.get("save_excel", True) and output_path:
        _log(log_callback, "Excel 결과표를 저장하는 중입니다...")
        saved_path = export_summary_xlsx(
            summary_layer,
            output_path,
            rows,
            total_count,
            total_area,
            report_style=options.get("report_style", True),
        )
        _log(log_callback, "Excel 저장 완료")

    return {
        "business_layer": business.name(),
        "cadastral_layer": cadastral.name(),
        "clipped_layer": clipped.name(),
        "summary_layer": summary_layer.name(),
        "category_count": len(rows),
        "total_count": total_count,
        "total_area_m2": total_area,
        "output_path": saved_path,
        "chat_lines": format_chat_table(
            rows,
            total_count,
            total_area,
        ),
    }
