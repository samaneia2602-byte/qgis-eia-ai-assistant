# -*- coding: utf-8 -*-

from collections import OrderedDict
from datetime import datetime


class ReportSection:
    """하나의 분석 결과를 공통 보고서 형식으로 보관합니다."""

    def __init__(
        self,
        section_id,
        title,
        columns,
        rows,
        summary=None,
        metadata=None,
    ):
        self.section_id = section_id
        self.title = title
        self.columns = list(columns or [])
        self.rows = list(rows or [])
        self.summary = dict(summary or {})
        self.metadata = dict(metadata or {})

    def to_dict(self):
        return {
            "section_id": self.section_id,
            "title": self.title,
            "columns": list(self.columns),
            "rows": list(self.rows),
            "summary": dict(self.summary),
            "metadata": dict(self.metadata),
        }


class ReportEngine:
    """
    지목·표고·경사·토지피복·생태자연도 등의 분석 결과를
    같은 구조로 모으는 공통 보고서 엔진입니다.
    """

    def __init__(
        self,
        project_title="환경영향평가 사업지역 분석",
        project_name="",
    ):
        self.project_title = project_title
        self.project_name = project_name
        self.created_at = datetime.now()
        self.sections = OrderedDict()
        self.metadata = {}

    def set_metadata(self, key, value):
        self.metadata[str(key)] = value

    def add_section(self, section):
        if not isinstance(section, ReportSection):
            raise TypeError("ReportSection 객체만 추가할 수 있습니다.")

        self.sections[section.section_id] = section
        return section

    def get_section(self, section_id):
        return self.sections.get(section_id)

    def has_section(self, section_id):
        return section_id in self.sections

    def to_dict(self):
        return {
            "project_title": self.project_title,
            "project_name": self.project_name,
            "created_at": self.created_at.isoformat(),
            "metadata": dict(self.metadata),
            "sections": [
                section.to_dict()
                for section in self.sections.values()
            ],
        }

    def build_text_summary(self):
        lines = [
            self.project_title,
            "=" * 48,
        ]

        if self.project_name:
            lines.append("사업명: %s" % self.project_name)

        lines.append(
            "작성일시: %s"
            % self.created_at.strftime("%Y-%m-%d %H:%M")
        )
        lines.append("")

        for section in self.sections.values():
            lines.append("[%s]" % section.title)

            for key, value in section.summary.items():
                lines.append("- %s: %s" % (key, value))

            lines.append("")

        return lines

    def export_excel(self, output_path):
        """
        openpyxl을 이용해 공통 보고서 Excel을 생성합니다.
        각 분석 항목은 별도 시트로 저장됩니다.
        """
        try:
            from openpyxl import Workbook
            from openpyxl.chart import PieChart, Reference
            from openpyxl.styles import (
                Alignment,
                Border,
                Font,
                PatternFill,
                Side,
            )
        except ImportError:
            raise RuntimeError(
                "보고서 Excel 생성을 위해 openpyxl이 필요합니다."
            )

        if not output_path.lower().endswith(".xlsx"):
            output_path += ".xlsx"

        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        thin = Side(style="thin", color="808080")
        border = Border(
            left=thin,
            right=thin,
            top=thin,
            bottom=thin,
        )
        header_fill = PatternFill(
            fill_type="solid",
            fgColor="D9E2F3",
        )
        total_fill = PatternFill(
            fill_type="solid",
            fgColor="E2F0D9",
        )

        for section in self.sections.values():
            sheet_name = self._safe_sheet_name(section.title)
            worksheet = workbook.create_sheet(sheet_name)

            column_count = max(1, len(section.columns))
            end_column = self._excel_column_name(column_count)

            worksheet.merge_cells(
                "A1:%s1" % end_column
            )
            worksheet["A1"] = section.title
            worksheet["A1"].font = Font(
                bold=True,
                size=14,
            )
            worksheet["A1"].alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

            row_index = 3
            for col_index, column in enumerate(
                section.columns,
                start=1,
            ):
                cell = worksheet.cell(
                    row=row_index,
                    column=col_index,
                    value=column,
                )
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(
                    horizontal="center"
                )
                cell.border = border

            data_start = row_index + 1

            for row in section.rows:
                row_index += 1
                for col_index, column in enumerate(
                    section.columns,
                    start=1,
                ):
                    value = row.get(column, "")
                    cell = worksheet.cell(
                        row=row_index,
                        column=col_index,
                        value=value,
                    )
                    cell.border = border

                    if isinstance(value, (int, float)):
                        if "면적(㎡)" in column:
                            cell.number_format = "#,##0.00"
                        elif "면적(ha)" in column:
                            cell.number_format = "0.0000"
                        elif "구성비" in column:
                            cell.number_format = "0.00"
                        elif isinstance(value, int):
                            cell.number_format = "#,##0"

                if str(row.get("지목", "")).strip() == "합계":
                    for cell in worksheet[row_index]:
                        cell.font = Font(bold=True)
                        cell.fill = total_fill

            summary_start = row_index + 2
            if section.summary:
                worksheet.cell(
                    row=summary_start,
                    column=1,
                    value="요약",
                ).font = Font(bold=True)

                for offset, (key, value) in enumerate(
                    section.summary.items(),
                    start=1,
                ):
                    worksheet.cell(
                        row=summary_start + offset,
                        column=1,
                        value=key,
                    )
                    worksheet.cell(
                        row=summary_start + offset,
                        column=2,
                        value=value,
                    )

            self._adjust_widths(
                worksheet,
                section.columns,
            )
            worksheet.freeze_panes = "A4"
            worksheet.sheet_view.showGridLines = False

            if (
                section.section_id == "cadastral"
                and len(section.rows) > 1
                and "지목" in section.columns
                and "면적(㎡)" in section.columns
            ):
                non_total_rows = [
                    row
                    for row in section.rows
                    if str(row.get("지목", "")).strip() != "합계"
                ]

                if non_total_rows:
                    category_col = (
                        section.columns.index("지목") + 1
                    )
                    value_col = (
                        section.columns.index("면적(㎡)") + 1
                    )
                    chart_end = (
                        data_start + len(non_total_rows) - 1
                    )

                    chart = PieChart()
                    chart.title = "지목별 면적 구성비"
                    labels = Reference(
                        worksheet,
                        min_col=category_col,
                        min_row=data_start,
                        max_row=chart_end,
                    )
                    data = Reference(
                        worksheet,
                        min_col=value_col,
                        min_row=data_start - 1,
                        max_row=chart_end,
                    )
                    chart.add_data(
                        data,
                        titles_from_data=True,
                    )
                    chart.set_categories(labels)
                    chart.height = 9
                    chart.width = 12
                    worksheet.add_chart(
                        chart,
                        "%s3" % self._excel_column_name(
                            column_count + 2
                        ),
                    )

        summary_sheet = workbook.create_sheet(
            "종합요약",
            0,
        )
        summary_sheet["A1"] = self.project_title
        summary_sheet["A1"].font = Font(
            bold=True,
            size=15,
        )
        summary_sheet["A3"] = "사업명"
        summary_sheet["B3"] = self.project_name
        summary_sheet["A4"] = "작성일시"
        summary_sheet["B4"] = self.created_at.strftime(
            "%Y-%m-%d %H:%M"
        )

        row_index = 6
        for section in self.sections.values():
            summary_sheet.cell(
                row=row_index,
                column=1,
                value=section.title,
            ).font = Font(bold=True)
            row_index += 1

            for key, value in section.summary.items():
                summary_sheet.cell(
                    row=row_index,
                    column=1,
                    value=key,
                )
                summary_sheet.cell(
                    row=row_index,
                    column=2,
                    value=value,
                )
                row_index += 1

            row_index += 1

        summary_sheet.column_dimensions["A"].width = 28
        summary_sheet.column_dimensions["B"].width = 28
        summary_sheet.sheet_view.showGridLines = False

        workbook.save(output_path)
        return output_path

    def _safe_sheet_name(self, title):
        invalid = set(r'[]:*?/\\')
        cleaned = "".join(
            "_" if char in invalid else char
            for char in title
        )
        return cleaned[:31] or "분석결과"

    def _adjust_widths(self, worksheet, columns):
        widths = {}
        for index, column in enumerate(
            columns,
            start=1,
        ):
            width = max(10, min(24, len(str(column)) * 2 + 4))
            widths[index] = width

        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                width = min(
                    36,
                    max(
                        widths.get(cell.column, 10),
                        len(str(cell.value)) + 3,
                    ),
                )
                widths[cell.column] = width

        for column_index, width in widths.items():
            worksheet.column_dimensions[
                self._excel_column_name(column_index)
            ].width = width

    def _excel_column_name(self, number):
        result = ""
        while number:
            number, remainder = divmod(number - 1, 26)
            result = chr(65 + remainder) + result
        return result


def build_cadastral_section(
    rows,
    total_count,
    total_area_m2,
    business_layer_name="",
    cadastral_layer_name="",
):
    report_rows = []

    for row in rows:
        report_rows.append(
            {
                "순번": len(report_rows) + 1,
                "지목": row["지목"],
                "필지수": int(row["필지수"]),
                "면적(㎡)": round(row["면적_m2"], 2),
                "면적(ha)": round(row["면적_ha"], 4),
                "구성비(%)": round(row["구성비_pct"], 2),
            }
        )

    report_rows.append(
        {
            "순번": len(report_rows) + 1,
            "지목": "합계",
            "필지수": int(total_count),
            "면적(㎡)": round(total_area_m2, 2),
            "면적(ha)": round(total_area_m2 / 10000.0, 4),
            "구성비(%)": 100.0,
        }
    )

    major = [
        row
        for row in rows
        if row.get("지목") != "미분류"
    ][:3]

    major_text = ", ".join(
        "%s %.2f%%"
        % (
            row["지목"],
            row["구성비_pct"],
        )
        for row in major
    )

    summary = {
        "총 필지수": "%s필지" % format(total_count, ","),
        "총면적": "%s㎡" % format(total_area_m2, ",.2f"),
        "총면적(ha)": "%.4fha" % (total_area_m2 / 10000.0),
        "주요 지목": major_text or "해당 없음",
    }

    metadata = {
        "사업지역 레이어": business_layer_name,
        "연속지적도 레이어": cadastral_layer_name,
    }

    return ReportSection(
        section_id="cadastral",
        title="사업지역 지목별 면적현황",
        columns=[
            "순번",
            "지목",
            "필지수",
            "면적(㎡)",
            "면적(ha)",
            "구성비(%)",
        ],
        rows=report_rows,
        summary=summary,
        metadata=metadata,
    )
